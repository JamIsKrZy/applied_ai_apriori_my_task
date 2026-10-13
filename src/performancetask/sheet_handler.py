import re
from collections.abc import Mapping, Sequence
from copy import copy
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.utils.cell import coordinate_to_tuple
from openpyxl.worksheet.table import Table, TableStyleInfo


def extract_table_entities(file_path, sheet_name, header_row, columns):
    """
    Extracts a floating table from an Excel sheet into a list of dictionaries.

    :param file_path: Path to the Excel file
    :param sheet_name: Name of the sheet
    :param header_row: The 0-indexed row number where the table headers are located
    :param columns: Excel column range (e.g., "C:F" if the table is in columns C through F)
    """
    # Read only the specific range containing the table
    df = pd.read_excel(
        file_path, sheet_name=sheet_name, header=header_row, usecols=columns
    )

    # Drop any completely empty rows that might exist below the table
    df.dropna(how="all", inplace=True)

    # Optional: Fill NaN values with None or empty strings if your entities require it
    df = df.where(pd.notnull(df), None)

    # Skip repeated header rows that appear within some source worksheets.
    entities = [
        entity
        for entity in df.to_dict(orient="records")
        if not all(entity.get(column) == column for column in df.columns)
    ]

    return entities


def _excel_value(value):
    if isinstance(value, (set, frozenset)):
        return ", ".join(str(item) for item in sorted(value, key=str))
    if isinstance(value, (list, tuple)):
        return ", ".join(str(item) for item in value)
    if isinstance(value, Mapping):
        return str(dict(value))
    return value


def _table_name(label: str, used_names: set[str]) -> str:
    base_name = re.sub(r"[^A-Za-z0-9_]", "_", label)
    if not base_name or not base_name[0].isalpha():
        base_name = f"Table_{base_name}"
    base_name = base_name[:200]

    name = base_name
    suffix = 2
    while name in used_names:
        name = f"{base_name[:195]}_{suffix}"
        suffix += 1
    used_names.add(name)
    return name


def export_results_to_excel(
    file_path: str | Path,
    tables: Mapping[str, Sequence[Mapping]],
    table_positions: Mapping[str, str] | None = None,
    sheet_name: str = "Apriori Results",
) -> None:
    """
    Replace the results sheet with one formatted Excel table per pipeline stage.

    `table_positions` maps a stage name to an Excel cell, such as
    `{"K-1 Itemsets": "G1"}`. Stages without an override stack vertically.
    """
    file_path = Path(file_path)
    file_path.parent.mkdir(parents=True, exist_ok=True)
    if file_path.exists():
        workbook = load_workbook(file_path)
    else:
        workbook = Workbook()
        workbook.remove(workbook.active)

    if sheet_name in workbook.sheetnames:
        del workbook[sheet_name]
    worksheet = workbook.create_sheet(sheet_name)
    positions = table_positions or {}
    unknown_positions = positions.keys() - tables.keys()
    if unknown_positions:
        raise ValueError(
            f"Table positions supplied for unknown stages: {', '.join(sorted(unknown_positions))}"
        )

    used_names: set[str] = set()
    occupied_ranges: list[tuple[int, int, int, int, str]] = []
    next_auto_row = 1

    for stage_name, records in tables.items():
        anchor = positions.get(stage_name, f"A{next_auto_row}")
        try:
            start_row, start_column = coordinate_to_tuple(anchor)
        except ValueError as error:
            raise ValueError(
                f"Invalid Excel position for {stage_name!r}: {anchor!r}"
            ) from error

        if start_row < 1 or start_column < 1:
            raise ValueError(f"Invalid Excel position for {stage_name!r}: {anchor!r}")

        records = list(records)
        headers = list(dict.fromkeys(key for record in records for key in record)) or [
            "Result"
        ]
        values = [
            [_excel_value(record.get(header)) for header in headers]
            for record in records
        ] or [["No records"]]

        title_row = start_row
        table_header_row = title_row + 1
        end_row = table_header_row + len(values)
        end_column = start_column + len(headers) - 1

        for (
            other_start_row,
            other_start_column,
            other_end_row,
            other_end_column,
            other_name,
        ) in occupied_ranges:
            overlaps = (
                start_row <= other_end_row
                and end_row >= other_start_row
                and start_column <= other_end_column
                and end_column >= other_start_column
            )
            if overlaps:
                raise ValueError(
                    f"Table {stage_name!r} overlaps {other_name!r}; "
                    "adjust table_positions."
                )
        occupied_ranges.append(
            (title_row, start_column, end_row, end_column, stage_name)
        )

        title_cell = worksheet.cell(
            row=title_row, column=start_column, value=stage_name
        )
        title_font = copy(title_cell.font)
        title_font.bold = True
        title_cell.font = title_font

        for offset, header in enumerate(headers):
            worksheet.cell(
                row=table_header_row,
                column=start_column + offset,
                value=str(header),
            )
        for row_offset, row_values in enumerate(values, start=1):
            for column_offset, value in enumerate(row_values):
                worksheet.cell(
                    row=table_header_row + row_offset,
                    column=start_column + column_offset,
                    value=value,
                )

        table_ref = (
            f"{worksheet.cell(row=table_header_row, column=start_column).coordinate}:"
            f"{worksheet.cell(row=end_row, column=end_column).coordinate}"
        )
        table = Table(
            displayName=_table_name(stage_name, used_names),
            ref=table_ref,
        )
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        worksheet.add_table(table)
        next_auto_row = max(next_auto_row, end_row + 3)

    workbook.save(file_path)


# Example Usage:
# If your table headers are on row 5 (index 4) and span from column C to F
# if __name__ == "__main__":
#     records = extract_table_entities(
#         file_path="data.xlsx",
#         sheet_name="Sheet1",
#         header_row=4,   # 0-indexed, so row 5 in Excel is index 4
#         columns="C:F"   # Only read columns C, D, E, and F
#     )

#     for entity in records:
#         print(entity)
